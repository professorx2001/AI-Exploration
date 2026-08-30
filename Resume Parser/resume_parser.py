import os
import json
from groq import Groq
from pathlib import Path
from docx import Document
from pypdf import PdfReader
from pydantic import BaseModel
from dotenv import load_dotenv


load_dotenv()
groq_api_key = os.getenv("GROQ_API_KEY")

if not groq_api_key:
    raise ValueError("Groq API Key not found!!!")

client = Groq(api_key=groq_api_key)

# Configuration
MODEL = "openai/gpt-oss-120b"
RESUME_DIR = Path(__file__).parent / "Resource"
SUPPORTED_FORMATS = {".pdf", ".docx"}

response_format = {
    "type": "json_object"
}

class JobDescription(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]

class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []

class MatchResult(BaseModel):
    score: float
    details: dict


job_description = """
Google Data Engineer
Minimum qualifications:
Bachelor's degree or equivalent practical experience.
4 years of experience in a data engineering, data infrastructure, or data analytics role.
Experience with database administration techniques or data engineering, as well as writing software in Java, C++, Python, Go, or JavaScript.

Preferred qualifications:
Masters degree in a relevant field.
5 years of experience in a data engineering, data infrastructure, or data analytics role.
Experience delivering and maintaining complex data projects from conception to production.
About the job

As a Data Engineer, you will take a significant role in designing and building the next generation of our data infrastructure. You will be responsible for architecting, implementing, and optimizing complex and scalable data pipelines, moving beyond basic development to own key components of our data warehouse. This role requires a strong technical expert who can manage massive datasets, write highly efficient SQL and Python code, and collaborate effectively with senior stakeholders and other engineers. You will not only build innovative data foundations, and AI-driven insights solutions, but also help define the standards and best practices that elevate the entire team, driving data quality and AI-readiness initiatives.

Responsibilities
Design, build, and maintain scalable data pipelines to ingest, process, and store data. Implement robust quality checks and monitoring to ensure data accuracy and reliability.
Write complex SQL queries for extraction, transformation, ad-hoc analysis, and automated reporting. Develop scalable data foundations and models designed to support AI/ML initiatives.
Develop, test, and deploy intelligent agents using Python and the Google ADK framework to automate tasks like data analysis, report generation, and system orchestration.
Partner with senior stakeholders, data scientists, and AI teams to understand complex requirements and architect robust long-term data solutions.
Develop tools to automate data processes, facilitate faster turnarounds, and increase efficiency. Monitor, troubleshoot, and tune data systems and pipelines.
"""



def parse_job_description():
    job_description_schema = JobDescription.model_json_schema()
    message_system = {
        "role": "system",
        "content": f"""
        You are an expert HR assistant.Your job is to analyze job descriptions and extract structured information from them.Return ONLY valid JSON matching this schema:{job_description_schema}
        IMPORTANT:
        Do NOT return the schema itself.
        Do NOT return fields like "properties", "title" or "type".
        Fill the schema with actual information extracted from the job description.
        If minimum experience is not mentioned, return null.
        If information for a list is missing, return an empty list.
        Do not invent information.
        """
    }
    
    message = {
        "role": "user",
        "content": f"Analyze the following job description:{job_description}"
    }

    messages = [message_system, message]

    response = client.chat.completions.create(
        model=MODEL, messages=messages, response_format=response_format)
    answer = response.choices[0].message.content

    job_description_raw_data = json.loads(answer)
    job_description_data = JobDescription(**job_description_raw_data)
    # JSON data of job description
    return job_description_data


def parse_resume_text(resume_text):
    """Parse resume text and extract structured information."""
    resume_schema = Resume.model_json_schema()
    message_system = {
        "role": "system",
        "content": f"""
        You are an expert resume parser.Extract information from the resume based on its meaning,not only based on exact section headings.
        Different resumes may use different headings.
        For example:
        - Experience
        - Professional Experience
        - Work History
        - Employment
        - Internships
        
        These may all contain relevant experience.
        Skills may also appear in the skills section, work experience,internships or projects.
        Return ONLY valid JSON matching this schema:{resume_schema}
        Important rules:
        1. Do not invent information.
        2. If a value is not available, return null.
        3. If a list has no information, return an empty list.
        4. Include internships inside experiences.
        5. Extract skills mentioned across the entire resume.
        """
    }
    message_user = {
        "role": "user",
        "content": f"Parse the following resume:{resume_text}"
    }
    messages = [message_system, message_user]
    response = client.chat.completions.create(model=MODEL, messages=messages, response_format=response_format)
    answer = response.choices[0].message.content
    resume_raw_data = json.loads(answer)
    parsed_resume = Resume(**resume_raw_data)

    return parsed_resume


def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def extract_text_from_document(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def extract_text_from_resume(file_path):
    """Extract text from PDF or DOCX files."""
    try:
        if file_path.suffix.lower() == ".pdf":
            return extract_text_from_pdf(file_path)
        elif file_path.suffix.lower() == ".docx":
            return extract_text_from_document(file_path)
        else:
            print(f"Unsupported file format: {file_path.suffix}")
            return None
    except Exception as e:
        print(f"Error extracting text from {file_path}: {e}")
        return None


def parse_resume_file(resume_file_path):
    """Parse a resume file and return structured resume data."""
    resume_text = extract_text_from_resume(resume_file_path)
    if not resume_text:
        return None
    return parse_resume_text(resume_text)


def match_resume_to_job_description(job_desc, resume):
    """Compare resume against job description and return match score."""
    match_result_schema = MatchResult.model_json_schema()

    message = {
        "role": "user",
        "content" : f"""
        You are an HR recruiter.Compare the candidate's resume with the job description.
        JOB DESCRIPTION: {job_desc}
        CANDIDATE RESUME:{resume}
        Return JSON matching this schema:{match_result_schema}
        Give me:
        1. Candidate name
        2. Matching skills
        3. Missing important skills
        4. Whether experience requirement is met
        5. Overall match percentage from 0 to 100
        6. A short final verdict
        Keep the response concise and easy to read.
        """
    }
    messages = [message]
    response = client.chat.completions.create(model=MODEL, messages=messages, response_format=response_format)
    answer = response.choices[0].message.content
    raw_data = json.loads(answer)
    data = MatchResult(**raw_data)

    return data


def get_resume_files(directory_path):
    """Get all valid resume files from directory."""
    if not directory_path.exists():
        print(f"Warning: Directory {directory_path} does not exist.")
        return []
    return [f for f in directory_path.iterdir() 
            if f.is_file() and f.suffix.lower() in SUPPORTED_FORMATS]


def process():
    """Main processing: parse JD, parse resumes, and match them."""
    job_desc = parse_job_description()
    results = []
    
    for resume_file in get_resume_files(Path(RESUME_DIR)):
        resume = parse_resume_file(resume_file)
        if resume:
            match = match_resume_to_job_description(job_desc, resume)
            results.append(match)
    
    return results

# python ./"Resume Parser"/resume_parser.py
print(process())